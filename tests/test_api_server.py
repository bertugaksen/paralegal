"""
Paralegal - Yerel REST API Sunucusu Entegrasyon Testleri
Frontend ile backend arasındaki tüm API uç noktalarını doğrular.
"""

import json
import os
import threading
import time
import urllib.request
import urllib.parse
from pathlib import Path
import pytest
from server import start_server

PORT = 52419
BASE_URL = f"http://127.0.0.1:{PORT}"

@pytest.fixture(scope="module")
def api_server():
    server = start_server(PORT)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    time.sleep(0.5)
    yield BASE_URL
    server.shutdown()


def post_json(url: str, data: dict) -> dict:
    req = urllib.request.Request(
        url,
        data=json.dumps(data).encode("utf-8"),
        headers={"Content-Type": "application/json"}
    )
    with urllib.request.urlopen(req) as resp:
        return json.loads(resp.read().decode("utf-8"))


def test_api_health(api_server):
    with urllib.request.urlopen(f"{api_server}/api/health") as resp:
        data = json.loads(resp.read().decode("utf-8"))
        assert data["status"] == "ok"
        assert data["offline"] is True


def test_api_load_and_scan_pdf(api_server, tmp_path):
    # Gerçek örnek belge
    sample_pdf = Path(__file__).parent.parent / "samples" / "ornekler" / "ornek_dava_dilekcesi.pdf"
    assert sample_pdf.exists()

    # 1. Load PDF
    meta = post_json(f"{api_server}/api/load-pdf", {"path": str(sample_pdf)})
    assert meta["total_pages"] >= 1
    assert meta["filename"] == "ornek_dava_dilekcesi.pdf"
    assert meta["is_scanned"] is False
    assert meta["has_text_layer"] is True
    assert meta["total_chars"] > 100

    # 2. Render Page Image
    img_url = f"{api_server}/api/pdf-page-image?path={urllib.parse.quote(str(sample_pdf))}&page=0&dpi=72"
    with urllib.request.urlopen(img_url) as resp:
        assert resp.status == 200
        assert resp.headers.get("Content-Type") == "image/jpeg"
        content = resp.read()
        assert len(content) > 1000

    # 3. Scan PDF
    scan_res = post_json(f"{api_server}/api/scan-pdf", {
        "path": str(sample_pdf),
        "filters": {
            "names": True,
            "companies": True,
            "dockets": True,
            "tckn": True,
            "iban": True,
            "phone": True,
            "plate": True,
            "email": True
        }
    })
    assert scan_res["total"] > 0
    assert scan_res["is_scanned"] is False
    assert scan_res["has_text_layer"] is True
    assert scan_res["total_chars"] > 100
    categories = {e["category"] for e in scan_res["entities"]}
    assert "TCKN" in categories or "KISI_ADI" in categories or "DOSYA_NO" in categories

    # 4. Search Text
    search_res = post_json(f"{api_server}/api/search-text", {
        "path": str(sample_pdf),
        "text": "Ahmet"
    })
    assert search_res["total"] >= 1

    # 5. Apply Redactions (kullanıcı tarafından onaylanan varlıklar)
    out_pdf = str(tmp_path / "redacted_output.pdf")
    entities_to_redact = [{**e, "approved": True} for e in scan_res["entities"][:3]]
    redact_res = post_json(f"{api_server}/api/apply-redactions", {
        "input_path": str(sample_pdf),
        "output_path": out_pdf,
        "redactions": entities_to_redact,
        "color": "black",
        "overlay": "none"
    })
    assert redact_res["success"] is True
    assert redact_res["redactions_applied"] >= 1
    assert os.path.exists(out_pdf)


def test_api_scanned_pdf_endpoints(api_server, tmp_path):
    """Taranmış (0 metinli görsel) PDF yüklendiğinde /api/load-pdf ve /api/scan-pdf meta verilerini doğrular."""
    import pymupdf as fitz
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 100, 100), 0)
    page.insert_image(page.rect, pixmap=pix)
    scanned_path = str(tmp_path / "api_taranmis_belge.pdf")
    doc.save(scanned_path)
    doc.close()

    # 1. /api/load-pdf
    load_meta = post_json(f"{api_server}/api/load-pdf", {"path": scanned_path})
    assert load_meta["total_pages"] == 1
    assert load_meta["is_scanned"] is True
    assert load_meta["has_text_layer"] is False
    assert load_meta["total_chars"] == 0

    # 2. /api/scan-pdf
    scan_meta = post_json(f"{api_server}/api/scan-pdf", {"path": scanned_path, "filters": {}})
    assert scan_meta["total"] == 0
    assert scan_meta["is_scanned"] is True
    assert scan_meta["has_text_layer"] is False
    assert scan_meta["total_chars"] == 0



def test_api_convert_images(api_server, tmp_path):
    sample_tiff = Path(__file__).parent.parent / "samples" / "ornekler" / "ornek_taranmis_delil.tiff"
    assert sample_tiff.exists()

    # 1. Thumbnail test
    thumb_url = f"{api_server}/api/image-thumbnail?path={urllib.parse.quote(str(sample_tiff))}&rotation=90"
    with urllib.request.urlopen(thumb_url) as resp:
        assert resp.status == 200
        assert resp.headers.get("Content-Type") == "image/jpeg"
        content = resp.read()
        assert len(content) > 100

    # 2. Convert with rotation test
    out_dir = str(tmp_path)
    conv_res = post_json(f"{api_server}/api/convert-images", {
        "input_paths": [str(sample_tiff)],
        "combine": False,
        "rotations": {str(sample_tiff): 90},
        "output_dir": out_dir
    })
    assert conv_res["success"] is True
    assert len(conv_res["output_paths"]) == 1
    assert os.path.exists(conv_res["output_paths"][0])


def test_api_convert_word(api_server, tmp_path):
    import docx
    sample_doc = str(tmp_path / "itiraz_dilekcesi.docx")
    d = docx.Document()
    d.add_heading("İSTANBUL 2. İCRA HUKUK MAHKEMESİ", level=1)
    d.add_paragraph("Borca ve yetkiye itirazlarımızın sunulmasıdır.")
    d.save(sample_doc)

    out_pdf = str(tmp_path / "itiraz_dilekcesi.pdf")
    res = post_json(f"{api_server}/api/convert-word", {
        "input_path": sample_doc,
        "output_path": out_pdf
    })
    assert res["success"] is True
    assert os.path.exists(out_pdf)


def test_api_hybrid_combine_word_and_image(api_server, tmp_path):
    import docx
    from PIL import Image

    doc_p = str(tmp_path / "dilekce_ek.docx")
    d = docx.Document()
    d.add_paragraph("Delil listesi ve açıklamaları.")
    d.save(doc_p)

    img_p = str(tmp_path / "delil_fotografi.png")
    im = Image.new("RGB", (300, 300), color=(180, 200, 220))
    im.save(img_p)

    conv_res = post_json(f"{api_server}/api/convert-images", {
        "input_paths": [doc_p, img_p],
        "combine": True,
        "output_dir": str(tmp_path)
    })
    assert conv_res["success"] is True
    assert len(conv_res["output_paths"]) == 1
    combined_pdf = conv_res["output_paths"][0]
    assert os.path.exists(combined_pdf)

    import pymupdf as fitz
    merged = fitz.open(combined_pdf)
    assert len(merged) >= 2
    merged.close()


def test_api_upload_file_base64_and_path(api_server, tmp_path):
    import base64
    sample_pdf = Path(__file__).parent.parent / "samples" / "ornekler" / "ornek_dava_dilekcesi.pdf"
    assert sample_pdf.exists()

    # 1. Existing local path passthrough
    res_path = post_json(f"{api_server}/api/upload-file", {"path": str(sample_pdf)})
    assert res_path["path"] == str(sample_pdf)
    assert res_path["filename"] == "ornek_dava_dilekcesi.pdf"

    # 2. Base64 file upload (browser drag & drop simulation)
    with open(sample_pdf, "rb") as f:
        b64_str = base64.b64encode(f.read()).decode("utf-8")

    res_upload = post_json(f"{api_server}/api/upload-file", {
        "filename": "dragged_dava.pdf",
        "data": f"data:application/pdf;base64,{b64_str}"
    })
    saved_path = res_upload["path"]
    assert os.path.exists(saved_path)
    assert saved_path.endswith(".pdf")
    assert "dragged_dava" in saved_path

    # Verify that the uploaded file can be loaded directly as a valid PDF
    meta = post_json(f"{api_server}/api/load-pdf", {"path": saved_path})
    assert meta["total_pages"] >= 1

def post_json_error(url: str, data: dict):
    """Hata donmesi beklenen istekler icin status code ve body dondurur."""
    import urllib.error
    req = urllib.request.Request(
        url,
        data=json.dumps(data).encode("utf-8"),
        headers={"Content-Type": "application/json"}
    )
    try:
        with urllib.request.urlopen(req) as resp:
            return resp.status, json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as err:
        return err.code, json.loads(err.read().decode("utf-8"))


def test_api_upload_file_different_types(api_server, tmp_path):
    """Farkli dosya turleri (Word, Excel, TIFF, PNG) icin upload ve entegrasyon dogrulamasi."""
    import base64
    from PIL import Image
    import docx
    import openpyxl

    # 1. WORD (.docx) TESTI
    word_file = str(tmp_path / "ornek_dilekce.docx")
    d = docx.Document()
    d.add_heading("ANKARA 1. ASLIYE HUKUK MAHKEMESI", level=1)
    d.add_paragraph("Faz P0 Drag & Drop Word Entegrasyon Testi")
    d.save(word_file)

    # Direct path
    res_word_path = post_json(f"{api_server}/api/upload-file", {"path": word_file})
    assert res_word_path["path"] == word_file
    assert res_word_path["filename"] == "ornek_dilekce.docx"

    # Base64 with data URL scheme
    with open(word_file, "rb") as f:
        word_b64 = base64.b64encode(f.read()).decode("utf-8")
    res_word_b64 = post_json(f"{api_server}/api/upload-file", {
        "filename": "mahkeme_dilekce_upload.docx",
        "data": f"data:application/vnd.openxmlformats-officedocument.wordprocessingml.document;base64,{word_b64}"
    })
    saved_word_path = res_word_b64["path"]
    assert os.path.exists(saved_word_path)
    assert saved_word_path.endswith(".docx")
    assert "mahkeme_dilekce_upload" in saved_word_path

    # Word to PDF conversion integration check
    word_pdf_out = str(tmp_path / "word_converted.pdf")
    conv_word = post_json(f"{api_server}/api/convert-word", {
        "input_path": saved_word_path,
        "output_path": word_pdf_out
    })
    assert conv_word["success"] is True
    assert os.path.exists(word_pdf_out)

    # 2. EXCEL (.xlsx) TESTI
    excel_file = str(tmp_path / "bilirkisi_hesaplama.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hesap Cetveli"
    ws.append(["Tarih", "Aciklama", "Tutar (TL)", "Faiz"])
    ws.append(["2026-01-01", "Kidem Tazminati", 150000, 12500])
    ws.append(["2026-02-01", "Ihbar Tazminati", 45000, 3200])
    wb.save(excel_file)

    # Direct path
    res_excel_path = post_json(f"{api_server}/api/upload-file", {"path": excel_file})
    assert res_excel_path["path"] == excel_file
    assert res_excel_path["filename"] == "bilirkisi_hesaplama.xlsx"

    # Base64 with content field name
    with open(excel_file, "rb") as f:
        excel_b64 = base64.b64encode(f.read()).decode("utf-8")
    res_excel_b64 = post_json(f"{api_server}/api/upload-file", {
        "name": "bilirkisi_upload.xlsx",
        "content": f"data:application/vnd.openxmlformats-officedocument.spreadsheetml.sheet;base64,{excel_b64}"
    })
    saved_excel_path = res_excel_b64["path"]
    assert os.path.exists(saved_excel_path)
    assert saved_excel_path.endswith(".xlsx")

    # Excel analysis and conversion integration check
    analysis = post_json(f"{api_server}/api/analyze-excel", {"path": saved_excel_path})
    assert len(analysis["sheets"]) >= 1
    assert "Hesap Cetveli" in analysis["sheets"]

    excel_pdf_out = str(tmp_path / "excel_converted.pdf")
    conv_excel = post_json(f"{api_server}/api/convert-excel", {
        "input_path": saved_excel_path,
        "output_path": excel_pdf_out
    })
    assert conv_excel["success"] is True
    assert os.path.exists(excel_pdf_out)

    # 3. TIFF GORSEL TESTI
    tiff_file = str(tmp_path / "taranmis_delil.tiff")
    im_tiff = Image.new("RGB", (200, 200), color=(240, 240, 240))
    im_tiff.save(tiff_file, format="TIFF")

    # Direct path
    res_tiff_path = post_json(f"{api_server}/api/upload-file", {"path": tiff_file})
    assert res_tiff_path["path"] == tiff_file
    assert res_tiff_path["filename"] == "taranmis_delil.tiff"

    # Base64 with base64 key
    with open(tiff_file, "rb") as f:
        tiff_b64 = base64.b64encode(f.read()).decode("utf-8")
    res_tiff_b64 = post_json(f"{api_server}/api/upload-file", {
        "filename": "delil_belgesi.tiff",
        "base64": tiff_b64  # raw base64 without data: prefix
    })
    saved_tiff_path = res_tiff_b64["path"]
    assert os.path.exists(saved_tiff_path)
    assert saved_tiff_path.endswith(".tiff")

    # Thumbnail check
    thumb_url = f"{api_server}/api/image-thumbnail?path={urllib.parse.quote(saved_tiff_path)}&rotation=0"
    with urllib.request.urlopen(thumb_url) as resp:
        assert resp.status == 200
        assert resp.headers.get("Content-Type") == "image/jpeg"

    # 4. PNG GORSEL TESTI
    png_file = str(tmp_path / "olay_yeri.png")
    im_png = Image.new("RGB", (150, 150), color=(100, 150, 200))
    im_png.save(png_file, format="PNG")

    with open(png_file, "rb") as f:
        png_b64 = base64.b64encode(f.read()).decode("utf-8")
    res_png = post_json(f"{api_server}/api/upload-file", {
        "filename": "olay_yeri_upload.png",
        "data": f"data:image/png;base64,{png_b64}"
    })
    saved_png_path = res_png["path"]
    assert os.path.exists(saved_png_path)
    assert saved_png_path.endswith(".png")

    # Combine convert test with uploaded TIFF and PNG
    combo_pdf = str(tmp_path / "gorseller_birlestirildi.pdf")
    combo_res = post_json(f"{api_server}/api/convert-images", {
        "input_paths": [saved_tiff_path, saved_png_path],
        "combine": True,
        "output_dir": str(tmp_path)
    })
    assert combo_res["success"] is True
    assert os.path.exists(combo_res["output_paths"][0])


def test_api_upload_file_security_and_unique_naming(api_server, tmp_path):
    """Guvenlik (path traversal engelleme), dosya cakismasi ve byte butunlugu testleri."""
    import base64
    import hashlib

    # 1. Path Traversal Engelleme
    content_raw = b"Gizli Hukuki Metin Icerigi"
    b64_data = base64.b64encode(content_raw).decode("utf-8")
    res_traversal = post_json(f"{api_server}/api/upload-file", {
        "filename": "../../windows/system32/kotu_amacli.pdf",
        "data": b64_data
    })
    saved_traversal_path = res_traversal["path"]
    assert os.path.exists(saved_traversal_path)
    # Paralegal uploads dizini disina cikamamali
    assert "paralegal_uploads" in saved_traversal_path
    assert "system32" not in saved_traversal_path.lower()
    assert Path(saved_traversal_path).name.startswith("kotu_amacli")

    # 2. Ayni isimle ard arda yuklemede benzersiz UUID ve uzerine yazmama garantisi
    res_1 = post_json(f"{api_server}/api/upload-file", {
        "filename": "vekaletname.pdf",
        "data": base64.b64encode(b"Vekaletname Surum 1").decode("utf-8")
    })
    res_2 = post_json(f"{api_server}/api/upload-file", {
        "filename": "vekaletname.pdf",
        "data": base64.b64encode(b"Vekaletname Surum 2").decode("utf-8")
    })
    assert res_1["path"] != res_2["path"]
    assert os.path.exists(res_1["path"])
    assert os.path.exists(res_2["path"])
    with open(res_1["path"], "rb") as f1, open(res_2["path"], "rb") as f2:
        assert f1.read() == b"Vekaletname Surum 1"
        assert f2.read() == b"Vekaletname Surum 2"

    # 3. Byte ve Hash Butunlugu
    sample_pdf = Path(__file__).parent.parent / "samples" / "ornekler" / "ornek_dava_dilekcesi.pdf"
    with open(sample_pdf, "rb") as f:
        original_bytes = f.read()
    orig_sha256 = hashlib.sha256(original_bytes).hexdigest()

    res_hash = post_json(f"{api_server}/api/upload-file", {
        "filename": "orijinal_kontrol.pdf",
        "data": base64.b64encode(original_bytes).decode("utf-8")
    })
    with open(res_hash["path"], "rb") as f:
        saved_bytes = f.read()
    assert hashlib.sha256(saved_bytes).hexdigest() == orig_sha256


def test_api_upload_file_error_handling(api_server):
    """Hata durumlari: Bos govde, gecersiz dosya yolu ve bozuk base64."""
    # 1. Bos govde ({}) -> 400 Bad Request
    status, body = post_json_error(f"{api_server}/api/upload-file", {})
    assert status == 400
    assert "Dosya" in body.get("error", "")

    # 2. Gecersiz/Varolmayan dosya yolu ve icerik yok -> 400 Bad Request
    status_path, body_path = post_json_error(f"{api_server}/api/upload-file", {
        "path": "C:/olmayan_klasor_12345/varolmayan_dosya.pdf"
    })
    assert status_path == 400
    assert "Dosya" in body_path.get("error", "")

    # 3. Bozuk Base64 verisi -> 500 Server Error ve sunucu cokmemesi
    status_corrupt, body_corrupt = post_json_error(f"{api_server}/api/upload-file", {
        "filename": "bozuk_dosya.pdf",
        "data": "!!!BuKesinlikleBase64Degildir==??"
    })
    assert status_corrupt == 500
    assert "error" in body_corrupt

    # 4. Sunucu hala saglikli calisiyor olmali
    with urllib.request.urlopen(f"{api_server}/api/health") as resp:
        health_data = json.loads(resp.read().decode("utf-8"))
        assert health_data["status"] == "ok"


def test_api_clean_temp_uploads(api_server):
    """Gecici yukleme dizini (%TEMP%/paralegal_uploads) temizleme dogrulamasi."""
    import tempfile
    import base64
    from core.routes import RouteHandler

    upload_dir = os.path.join(tempfile.gettempdir(), "paralegal_uploads")
    os.makedirs(upload_dir, exist_ok=True)

    # Test amacli bir yukleme yap
    res = post_json(f"{api_server}/api/upload-file", {
        "filename": "temizlenecek_belge.pdf",
        "data": base64.b64encode(b"Temizleme testi icerigi").decode("utf-8")
    })
    assert os.path.exists(res["path"])

    # API uzerinden temizleme istegi gonder (max_age_seconds: 0 ile tumunu sil)
    clean_res = post_json(f"{api_server}/api/clean-temp-uploads", {"max_age_seconds": 0})
    assert clean_res["success"] is True
    assert clean_res["deleted_count"] >= 1
    assert not os.path.exists(res["path"])

    # RouteHandler statik metodu dogrudan cagrisi
    dummy_file = os.path.join(upload_dir, "statik_temizleme_testi.txt")
    with open(dummy_file, "w") as f:
        f.write("test")
    assert os.path.exists(dummy_file)
    count = RouteHandler.cleanup_temp_uploads(0)
    assert count >= 1
    assert not os.path.exists(dummy_file)


def test_api_file_sizes(api_server, tmp_path):
    """Dosya boyutu hesaplama API uç noktasını doğrular."""
    # 1. Test dosyaları oluştur
    f1 = tmp_path / "test_doc1.pdf"
    f1.write_bytes(b"0" * 1024 * 500)  # ~500 KB

    f2 = tmp_path / "test_doc2.pdf"
    f2.write_bytes(b"0" * (1024 * 1024 * 2 + 512 * 1024))  # ~2.5 MB

    non_existent = str(tmp_path / "non_existent.pdf")

    # 2. API çağrısı
    res = post_json(f"{api_server}/api/file-sizes", {
        "paths": [str(f1), str(f2), non_existent]
    })

    assert isinstance(res, list)
    assert len(res) == 3

    assert res[0]["path"] == str(f1)
    assert res[0]["size_bytes"] == 500 * 1024
    assert "KB" in res[0]["size_formatted"]

    assert res[1]["path"] == str(f2)
    assert res[1]["size_bytes"] == 1024 * 1024 * 2 + 512 * 1024
    assert "MB" in res[1]["size_formatted"]

    assert res[2]["path"] == non_existent
    assert res[2]["size_bytes"] == 0


