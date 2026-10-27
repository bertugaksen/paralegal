"""
Tests for Paralegal Word Converter Module (.docx and .doc to PDF)
Includes MS Word COM automation unit and mock tests, cascading pipeline tests,
LibreOffice fallback, and ReportLab pure Python fallback.
"""

import os
import sys
import shutil
import tempfile
from unittest.mock import MagicMock, patch
import pytest
import pymupdf as fitz
from pathlib import Path

from core.word_converter import (
    WordConverter,
    HAS_WIN32COM,
    HAS_PYTHON_DOCX,
    HAS_REPORTLAB,
    WD_FORMAT_PDF,
    WD_EXPORT_FORMAT_PDF,
    WD_DO_NOT_SAVE_CHANGES,
    WD_ALERTS_NONE,
)


class TestWordConverter:
    """Test suite for Word (.docx, .doc) conversion engine and cascading pipeline."""

    @pytest.fixture
    def sample_docx_file(self, tmp_path):
        """Creates a genuine legal petition .docx file with title, body, and table."""
        import docx

        doc_path = str(tmp_path / "dava_dilekcesi.docx")
        doc = docx.Document()
        doc.add_heading("İSTANBUL ANADOLU 4. ASLİYE TİCARET MAHKEMESİ'NE", level=1)
        doc.add_paragraph("DAVACI: Anadolu Lojistik A.Ş.")
        doc.add_paragraph("VEKİLİ: Av. Mehmet Demir")
        doc.add_paragraph("DAVALI: Marmara İnşaat San. Tic. Ltd. Şti.")
        doc.add_paragraph("KONU: 150.000,00 TL ticari alacağın tahsili talebidir.")

        # Add a table
        table = doc.add_table(rows=3, cols=3)
        headers = ["Fatura No", "Fatura Tarihi", "Tutar (TL)"]
        for i, h in enumerate(headers):
            table.cell(0, i).text = h

        row1 = ["FT-2024-001", "12.01.2024", "75.000,00"]
        for i, val in enumerate(row1):
            table.cell(1, i).text = val

        row2 = ["FT-2024-002", "18.02.2024", "75.000,00"]
        for i, val in enumerate(row2):
            table.cell(2, i).text = val

        doc.save(doc_path)
        return doc_path

    @pytest.fixture
    def sample_doc_file(self, tmp_path):
        """Creates a legacy Word (.doc) compatible file."""
        doc_path = str(tmp_path / "eski_vekaletname.doc")
        with open(doc_path, "w", encoding="utf-8") as f:
            f.write(
                r"{\rtf1\ansi\deff0 {\fonttbl {\f0 Times New Roman;}} \f0\fs24 GENEL VEKALETNAME \par Ahzu kabz, sulh ve ibraya mezunum.}"
            )
        return doc_path

    @pytest.mark.parametrize(
        "filename, expected",
        [
            ("dilekce.docx", True),
            ("DILEKCE.DOCX", True),
            ("vekaletname.doc", True),
            ("ESKI_EVRAK.DOC", True),
            ("rapor.pdf", False),
            ("hesap.xlsx", False),
            ("resim.png", False),
            ("notlar.txt", False),
            ("isimsiz", False),
        ],
    )
    def test_is_supported(self, filename, expected):
        assert WordConverter.is_supported(filename) == expected

    def test_find_libreoffice(self):
        soffice = WordConverter.find_libreoffice()
        if soffice is not None:
            assert os.path.exists(soffice)
        else:
            assert soffice is None

    # --- MS Word COM Availability Unit Tests ---

    def test_is_ms_word_available_detected(self, monkeypatch):
        mock_key = MagicMock()
        mock_winreg = MagicMock()
        mock_winreg.OpenKey.return_value = mock_key

        monkeypatch.setattr("sys.platform", "win32")
        monkeypatch.setattr("core.word_converter.HAS_WIN32COM", True)
        monkeypatch.setitem(sys.modules, "winreg", mock_winreg)

        assert WordConverter.is_ms_word_available() is True
        mock_key.Close.assert_called_once()

    def test_is_ms_word_available_missing_registry(self, monkeypatch):
        mock_winreg = MagicMock()
        mock_winreg.OpenKey.side_effect = FileNotFoundError("Key not found")

        monkeypatch.setattr("sys.platform", "win32")
        monkeypatch.setattr("core.word_converter.HAS_WIN32COM", True)
        monkeypatch.setitem(sys.modules, "winreg", mock_winreg)

        assert WordConverter.is_ms_word_available() is False

    def test_is_ms_word_available_non_windows(self, monkeypatch):
        monkeypatch.setattr("sys.platform", "linux")
        assert WordConverter.is_ms_word_available() is False

    def test_is_ms_word_available_without_win32com(self, monkeypatch):
        monkeypatch.setattr("sys.platform", "win32")
        monkeypatch.setattr("core.word_converter.HAS_WIN32COM", False)
        assert WordConverter.is_ms_word_available() is False

    # --- MS Word COM Automation Unit & Mock Tests ---

    def test_convert_via_ms_word_success(self, sample_docx_file, tmp_path, monkeypatch):
        out_pdf = str(tmp_path / "com_out.pdf")

        mock_word = MagicMock()
        mock_doc = MagicMock()
        mock_word.Documents.Open.return_value = mock_doc

        def fake_export(OutputFileName, ExportFormat, **kwargs):
            doc = fitz.open()
            page = doc.new_page()
            page.insert_text((50, 50), "MS Word Export Success")
            doc.save(OutputFileName)
            doc.close()

        mock_doc.ExportAsFixedFormat.side_effect = fake_export

        mock_win32com = MagicMock()
        mock_win32com.client.DispatchEx.return_value = mock_word

        mock_pythoncom = MagicMock()

        monkeypatch.setattr("sys.platform", "win32")
        monkeypatch.setattr("core.word_converter.HAS_WIN32COM", True)
        monkeypatch.setattr("core.word_converter.win32com", mock_win32com, raising=False)
        monkeypatch.setattr("core.word_converter.pythoncom", mock_pythoncom, raising=False)

        res = WordConverter._convert_via_ms_word(sample_docx_file, out_pdf)

        assert os.path.exists(res)
        assert res == os.path.abspath(out_pdf)

        # COM lifecycle verifications
        mock_pythoncom.CoInitialize.assert_called_once()
        mock_pythoncom.CoUninitialize.assert_called_once()
        assert mock_word.Visible is False
        assert mock_word.DisplayAlerts == WD_ALERTS_NONE
        mock_word.Documents.Open.assert_called_once()
        mock_doc.ExportAsFixedFormat.assert_called_once()
        mock_doc.Close.assert_called_once_with(WD_DO_NOT_SAVE_CHANGES)
        mock_word.Quit.assert_called_once_with(WD_DO_NOT_SAVE_CHANGES)

    def test_convert_via_ms_word_saveas_fallback(self, sample_docx_file, tmp_path, monkeypatch):
        out_pdf = str(tmp_path / "com_saveas_out.pdf")

        mock_word = MagicMock()
        mock_doc = MagicMock()
        mock_word.Documents.Open.return_value = mock_doc

        # ExportAsFixedFormat fails
        mock_doc.ExportAsFixedFormat.side_effect = Exception("ExportAsFixedFormat not supported")

        def fake_saveas(FileName, FileFormat, **kwargs):
            doc = fitz.open()
            page = doc.new_page()
            page.insert_text((50, 50), "MS Word SaveAs Success")
            doc.save(FileName)
            doc.close()

        mock_doc.SaveAs.side_effect = fake_saveas

        mock_win32com = MagicMock()
        mock_win32com.client.DispatchEx.return_value = mock_word
        mock_pythoncom = MagicMock()

        monkeypatch.setattr("sys.platform", "win32")
        monkeypatch.setattr("core.word_converter.HAS_WIN32COM", True)
        monkeypatch.setattr("core.word_converter.win32com", mock_win32com, raising=False)
        monkeypatch.setattr("core.word_converter.pythoncom", mock_pythoncom, raising=False)

        res = WordConverter._convert_via_ms_word(sample_docx_file, out_pdf)

        assert os.path.exists(res)
        mock_doc.SaveAs.assert_called_once()
        mock_doc.Close.assert_called_once_with(WD_DO_NOT_SAVE_CHANGES)
        mock_word.Quit.assert_called_once_with(WD_DO_NOT_SAVE_CHANGES)

    def test_convert_via_ms_word_cleanup_on_error(self, sample_docx_file, tmp_path, monkeypatch):
        out_pdf = str(tmp_path / "com_err_out.pdf")

        mock_word = MagicMock()
        mock_doc = MagicMock()
        mock_word.Documents.Open.return_value = mock_doc

        mock_doc.ExportAsFixedFormat.side_effect = RuntimeError("Fatal Word Error")
        mock_doc.SaveAs.side_effect = RuntimeError("SaveAs Fatal Error")

        mock_win32com = MagicMock()
        mock_win32com.client.DispatchEx.return_value = mock_word
        mock_pythoncom = MagicMock()

        monkeypatch.setattr("sys.platform", "win32")
        monkeypatch.setattr("core.word_converter.HAS_WIN32COM", True)
        monkeypatch.setattr("core.word_converter.win32com", mock_win32com, raising=False)
        monkeypatch.setattr("core.word_converter.pythoncom", mock_pythoncom, raising=False)

        with pytest.raises(RuntimeError):
            WordConverter._convert_via_ms_word(sample_docx_file, out_pdf)

        # Cleanup must still occur even after fatal errors
        mock_doc.Close.assert_called_once_with(WD_DO_NOT_SAVE_CHANGES)
        mock_word.Quit.assert_called_once_with(WD_DO_NOT_SAVE_CHANGES)
        mock_pythoncom.CoUninitialize.assert_called_once()

    def test_convert_via_ms_word_dispatch_fallback(self, sample_docx_file, tmp_path, monkeypatch):
        out_pdf = str(tmp_path / "com_dispatch_fallback.pdf")

        mock_word = MagicMock()
        mock_doc = MagicMock()
        mock_word.Documents.Open.return_value = mock_doc

        def fake_export(OutputFileName, ExportFormat, **kwargs):
            doc = fitz.open()
            page = doc.new_page()
            page.insert_text((50, 50), "Dispatch Fallback Text")
            doc.save(OutputFileName)
            doc.close()

        mock_doc.ExportAsFixedFormat.side_effect = fake_export

        mock_win32com = MagicMock()
        mock_win32com.client.DispatchEx.side_effect = Exception("DispatchEx failed")
        mock_win32com.client.Dispatch.return_value = mock_word
        mock_pythoncom = MagicMock()

        monkeypatch.setattr("sys.platform", "win32")
        monkeypatch.setattr("core.word_converter.HAS_WIN32COM", True)
        monkeypatch.setattr("core.word_converter.win32com", mock_win32com, raising=False)
        monkeypatch.setattr("core.word_converter.pythoncom", mock_pythoncom, raising=False)

        res = WordConverter._convert_via_ms_word(sample_docx_file, out_pdf)

        assert os.path.exists(res)
        mock_win32com.client.Dispatch.assert_called_once_with("Word.Application")

    # --- Cascading Pipeline Tests ---

    def test_cascading_pipeline_word_priority(self, sample_docx_file, tmp_path, monkeypatch):
        """Verifies 1st priority: MS Word is called before LibreOffice or ReportLab."""
        out_pdf = str(tmp_path / "priority1.pdf")
        called = []

        def mock_word_convert(cls, in_p, out_p):
            called.append("ms_word")
            doc = fitz.open()
            page = doc.new_page()
            page.insert_text((50, 50), "Priority 1 Word")
            doc.save(out_p)
            doc.close()
            return out_p

        monkeypatch.setattr(WordConverter, "is_ms_word_available", classmethod(lambda cls: True))
        monkeypatch.setattr(WordConverter, "_convert_via_ms_word", classmethod(mock_word_convert))
        monkeypatch.setattr(
            WordConverter,
            "find_libreoffice",
            classmethod(lambda cls: called.append("libreoffice") or "soffice.exe"),
        )

        res = WordConverter.convert_word_to_pdf(sample_docx_file, out_pdf)
        assert os.path.exists(res)
        assert called == ["ms_word"]  # LibreOffice should NOT be touched when MS Word succeeds

    def test_cascading_pipeline_word_fails_to_libreoffice(
        self, sample_docx_file, tmp_path, monkeypatch
    ):
        """Verifies 2nd priority: When MS Word fails, LibreOffice is called."""
        out_pdf = str(tmp_path / "priority2.pdf")
        stages = []

        def mock_word_fail(cls, in_p, out_p):
            stages.append("ms_word_failed")
            raise RuntimeError("Word crashed")

        monkeypatch.setattr(WordConverter, "is_ms_word_available", classmethod(lambda cls: True))
        monkeypatch.setattr(WordConverter, "_convert_via_ms_word", classmethod(mock_word_fail))
        monkeypatch.setattr(
            WordConverter, "find_libreoffice", classmethod(lambda cls: "mock_soffice.exe")
        )

        # Mock subprocess.run for LibreOffice
        def mock_run(cmd, **kwargs):
            stages.append("libreoffice_called")
            out_dir = cmd[cmd.index("--outdir") + 1]
            gen_pdf = os.path.join(out_dir, Path(sample_docx_file).stem + ".pdf")
            doc = fitz.open()
            page = doc.new_page()
            page.insert_text((50, 50), "LibreOffice Output")
            doc.save(gen_pdf)
            doc.close()
            res_obj = MagicMock()
            res_obj.returncode = 0
            return res_obj

        monkeypatch.setattr("subprocess.run", mock_run)

        res = WordConverter.convert_word_to_pdf(sample_docx_file, out_pdf)
        assert os.path.exists(res)
        assert "ms_word_failed" in stages
        assert "libreoffice_called" in stages

    def test_cascading_pipeline_word_and_libreoffice_fail_to_reportlab(
        self, sample_docx_file, tmp_path, monkeypatch
    ):
        """Verifies 3rd priority: When both Word and LibreOffice fail, ReportLab handles .docx."""
        out_pdf = str(tmp_path / "priority3.pdf")
        stages = []

        def mock_word_fail(cls, in_p, out_p):
            stages.append("word_failed")
            raise RuntimeError("Word not usable")

        monkeypatch.setattr(WordConverter, "is_ms_word_available", classmethod(lambda cls: True))
        monkeypatch.setattr(WordConverter, "_convert_via_ms_word", classmethod(mock_word_fail))
        monkeypatch.setattr(WordConverter, "find_libreoffice", classmethod(lambda cls: None))

        res = WordConverter.convert_word_to_pdf(sample_docx_file, out_pdf)
        assert os.path.exists(res)
        assert "word_failed" in stages

        doc = fitz.open(res)
        text = doc[0].get_text()
        assert "TİCARET MAHKEMESİ" in text or "DAVACI" in text
        doc.close()

    def test_convert_docx_to_pdf_integration(self, sample_docx_file, tmp_path):
        """Integration test for convert_word_to_pdf with cascading pipeline."""
        out_pdf = str(tmp_path / "dava_dilekcesi_entegrasyon.pdf")
        callbacks = []

        def on_prog(curr, tot, msg):
            callbacks.append((curr, tot, msg))

        res_path = WordConverter.convert_word_to_pdf(
            sample_docx_file, output_path=out_pdf, progress_callback=on_prog
        )

        assert os.path.exists(res_path)
        assert res_path == out_pdf
        assert len(callbacks) > 0

        # Verify PDF structure using PyMuPDF
        doc = fitz.open(res_path)
        assert len(doc) >= 1
        page_text = doc[0].get_text()
        assert "TİCARET MAHKEMESİ" in page_text or "DAVACI" in page_text or "Anadolu" in page_text
        doc.close()

    def test_convert_doc_to_pdf(self, sample_doc_file, tmp_path, monkeypatch):
        """Tests .doc conversion via MS Word COM (mocked)."""
        out_pdf = str(tmp_path / "eski_vekaletname.pdf")

        def mock_ms_word(cls, input_p, output_p):
            doc = fitz.open()
            page = doc.new_page()
            page.insert_text((50, 50), "GENEL VEKALETNAME Ahzu kabz, sulh ve ibraya mezunum.")
            doc.save(output_p)
            doc.close()
            return output_p

        monkeypatch.setattr(WordConverter, "is_ms_word_available", classmethod(lambda cls: True))
        monkeypatch.setattr(WordConverter, "_convert_via_ms_word", classmethod(mock_ms_word))

        res_path = WordConverter.convert_word_to_pdf(sample_doc_file, output_path=out_pdf)

        assert os.path.exists(res_path)
        doc = fitz.open(res_path)
        assert len(doc) >= 1
        page_text = doc[0].get_text()
        assert "VEKALETNAME" in page_text or "sulh" in page_text
        doc.close()

    def test_convert_doc_fails_gracefully_when_no_engines(
        self, sample_doc_file, tmp_path, monkeypatch
    ):
        """When MS Word and LibreOffice are absent, .doc raises descriptive RuntimeError."""
        monkeypatch.setattr(WordConverter, "is_ms_word_available", classmethod(lambda cls: False))
        monkeypatch.setattr(WordConverter, "find_libreoffice", classmethod(lambda cls: None))

        out_pdf = str(tmp_path / "doc_fail.pdf")
        with pytest.raises(RuntimeError, match="Word dönüştürme gerçekleştirilemedi"):
            WordConverter.convert_word_to_pdf(sample_doc_file, output_path=out_pdf)

    def test_reportlab_fallback_direct(self, sample_docx_file, tmp_path):
        if not (HAS_PYTHON_DOCX and HAS_REPORTLAB):
            pytest.skip("python-docx veya reportlab yüklü değil")

        fallback_pdf = str(tmp_path / "fallback_out.pdf")
        res_path = WordConverter._convert_via_reportlab(sample_docx_file, fallback_pdf)

        assert os.path.exists(res_path)
        doc = fitz.open(res_path)
        assert len(doc) >= 1
        page_text = doc[0].get_text()
        assert "DAVACI" in page_text or "TİCARET" in page_text
        doc.close()

    def test_convert_nonexistent_file_raises(self, tmp_path):
        nonexistent = str(tmp_path / "yok.docx")
        with pytest.raises(FileNotFoundError, match="Word dosyası bulunamadı"):
            WordConverter.convert_word_to_pdf(nonexistent)

    def test_convert_unsupported_format_raises(self, tmp_path):
        dummy_txt = str(tmp_path / "test.xyz")
        Path(dummy_txt).write_text("dummy")
        with pytest.raises(ValueError, match="Desteklenmeyen dosya formatı"):
            WordConverter.convert_word_to_pdf(dummy_txt)

    def test_uyap_10mb_protection_applied(self, sample_docx_file, tmp_path):
        out_pdf = str(tmp_path / "uyap_guvenli.pdf")
        res_path = WordConverter.convert_word_to_pdf(sample_docx_file, output_path=out_pdf)
        assert os.path.exists(res_path)
        assert os.path.getsize(res_path) <= 10 * 1024 * 1024
