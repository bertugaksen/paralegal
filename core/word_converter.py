"""
Paralegal - High-Fidelity Word (.docx, .doc) to UYAP PDF Converter Module
%100 Yerel, Sıfır İnternet, Sıfır Telemetri.

Desteklenen formatlar: .docx, .doc (Dava/cevap dilekçeleri, sözleşmeler, ihtarnameler, müzekkereler).
Kademeli Mimari (Cascading Pipeline):
1. Öncelik 1: Microsoft Office Word (Windows COM - win32com.client: %100 aslına uygun mizanpaj, antet, imza blokları ve fontlar).
2. Öncelik 2: LibreOffice Headless (Tam aslına uygun mizanpaj, yazı tipleri, kenar boşlukları ve tablolar).
3. Öncelik 3: Saf Python Fallback (python-docx + ReportLab - Harici ofis programı bulunmadığında çalışan yerel kurtarma motoru).
4. UYAP 10 MB Otomatik Koruma: Üretilen PDF 10 MB sınırını aşıyorsa kalite kaybı yaşanmadan optimize edilir.
"""

import os
import sys
import shutil
import tempfile
import subprocess
from pathlib import Path
from typing import Optional, Callable

try:
    import pythoncom
    import win32com.client
    HAS_WIN32COM = True
except ImportError:
    HAS_WIN32COM = False

try:
    import docx
    HAS_PYTHON_DOCX = True
except ImportError:
    HAS_PYTHON_DOCX = False

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    HAS_REPORTLAB = True
except ImportError:
    HAS_REPORTLAB = False


# Word COM Constants
WD_FORMAT_PDF = 17
WD_EXPORT_FORMAT_PDF = 17
WD_DO_NOT_SAVE_CHANGES = 0
WD_ALERTS_NONE = 0


class WordConverter:
    """Word (.docx, .doc) belgelerini aslına uygun ve UYAP uyumlu PDF'e dönüştüren motor."""

    SUPPORTED_EXTENSIONS = {".docx", ".doc"}

    # Word COM Constants
    WD_FORMAT_PDF = WD_FORMAT_PDF
    WD_EXPORT_FORMAT_PDF = WD_EXPORT_FORMAT_PDF
    WD_DO_NOT_SAVE_CHANGES = WD_DO_NOT_SAVE_CHANGES
    WD_ALERTS_NONE = WD_ALERTS_NONE

    @classmethod
    def is_supported(cls, file_path: str) -> bool:
        """Dosyanın desteklenen bir Word formatı olup olmadığını denetler."""
        return Path(file_path).suffix.lower() in cls.SUPPORTED_EXTENSIONS

    @classmethod
    def is_ms_word_available(cls) -> bool:
        """Windows ortamında Microsoft Word COM otomasyonunun kullanılabilir olup olmadığını denetler."""
        if sys.platform != "win32" or not HAS_WIN32COM:
            return False
        try:
            import winreg
            key = winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, "Word.Application")
            key.Close()
            return True
        except Exception:
            return False

    @classmethod
    def find_libreoffice(cls) -> Optional[str]:
        """Sistemde kurulu LibreOffice (soffice) çalıştırılabilir dosyasını bulur."""
        # 1. PATH kontrolü
        for cmd_name in ("soffice", "libreoffice"):
            found = shutil.which(cmd_name)
            if found and os.path.exists(found):
                return found

        # 2. Windows Standart Kurulum Yolları
        windows_paths = [
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
            r"C:\Program Files\LibreOffice\program\soffice.com",
            os.path.expandvars(r"%LOCALAPPDATA%\Programs\LibreOffice\program\soffice.exe"),
        ]
        for p in windows_paths:
            if os.path.exists(p):
                return p

        return None

    @classmethod
    def _convert_via_ms_word(cls, input_path: str, output_path: str) -> str:
        """
        Microsoft Word COM otomasyonu (win32com.client) kullanarak
        belgeyi %100 aslına uygun şekilde PDF'e dönüştürür.
        """
        if not (sys.platform == "win32" and HAS_WIN32COM):
            raise RuntimeError("Windows COM ve win32com.client mevcut değil.")

        abs_in = os.path.abspath(input_path)
        abs_out = os.path.abspath(output_path)

        out_dir = os.path.dirname(abs_out)
        if out_dir:
            os.makedirs(out_dir, exist_ok=True)

        # COM kütüphanesini bu iş parçacığı için başlat (Multi-threading desteği)
        pythoncom.CoInitialize()

        word = None
        doc = None
        try:
            # Bağımsız yeni bir Word örneği oluştur
            try:
                word = win32com.client.DispatchEx("Word.Application")
            except Exception:
                word = win32com.client.Dispatch("Word.Application")

            word.Visible = False
            word.DisplayAlerts = cls.WD_ALERTS_NONE

            # Belgeyi salt okunur olarak aç
            doc = word.Documents.Open(
                FileName=abs_in,
                ReadOnly=True,
                ConfirmConversions=False,
                AddToRecentFiles=False,
            )

            # 1. Tercih: ExportAsFixedFormat (Word 2007 SP2+ / 2010+)
            converted = False
            try:
                try:
                    doc.ExportAsFixedFormat(
                        OutputFileName=abs_out,
                        ExportFormat=cls.WD_EXPORT_FORMAT_PDF,
                        OpenAfterExport=False,
                        OptimizeFor=0,     # wdExportOptimizeForPrint
                        CreateBookmarks=1, # wdExportCreateHeadingBookmarks
                        DocStructureTags=True,
                    )
                except TypeError:
                    doc.ExportAsFixedFormat(abs_out, cls.WD_EXPORT_FORMAT_PDF)
                converted = True
            except Exception:
                converted = False

            # 2. Tercih: SaveAs (wdFormatPDF = 17)
            if not converted or not os.path.exists(abs_out):
                try:
                    doc.SaveAs(FileName=abs_out, FileFormat=cls.WD_FORMAT_PDF)
                except TypeError:
                    doc.SaveAs(abs_out, cls.WD_FORMAT_PDF)

            if not os.path.exists(abs_out) or os.path.getsize(abs_out) == 0:
                raise RuntimeError(f"Word COM PDF üretemedi: {abs_out}")

            return abs_out
        finally:
            was_faulthandler_enabled = False
            try:
                import faulthandler
                was_faulthandler_enabled = faulthandler.is_enabled()
                if was_faulthandler_enabled:
                    faulthandler.disable()
            except Exception:
                pass

            try:
                if doc is not None:
                    try:
                        doc.Close(cls.WD_DO_NOT_SAVE_CHANGES)
                    except Exception:
                        pass
                    doc = None
                if word is not None:
                    try:
                        word.Quit(cls.WD_DO_NOT_SAVE_CHANGES)
                    except Exception:
                        pass
                    word = None
                import gc
                gc.collect()
                try:
                    pythoncom.CoUninitialize()
                except Exception:
                    pass
            finally:
                if was_faulthandler_enabled:
                    try:
                        faulthandler.enable()
                    except Exception:
                        pass

    @classmethod
    def convert_word_to_pdf(
        cls,
        input_path: str,
        output_path: Optional[str] = None,
        progress_callback: Optional[Callable[[int, int, str], None]] = None,
    ) -> str:
        """
        Word (.docx / .doc) belgesini adli standartlara uygun olarak PDF formatına dönüştürür.

        Kademeli Öncelik Sıralaması (Cascading Pipeline):
        1. Öncelik: Microsoft Office Word (Windows COM - %100 Aslına Uygunluk)
        2. Öncelik: LibreOffice Headless (soffice)
        3. Öncelik: Saf Python Fallback (python-docx + ReportLab)
        """
        in_p = Path(input_path)
        if not in_p.exists():
            raise FileNotFoundError(f"Word dosyası bulunamadı: {input_path}")

        if not cls.is_supported(str(in_p)):
            raise ValueError(f"Desteklenmeyen dosya formatı: {in_p.suffix}")

        if output_path is None:
            output_path = str(in_p.with_suffix(".pdf"))
        else:
            out_dir = os.path.dirname(output_path)
            if out_dir:
                os.makedirs(out_dir, exist_ok=True)

        if progress_callback:
            progress_callback(1, 3, f"Word belgesi taranıyor: {in_p.name}")

        # 1. ÖNCELİK: Microsoft Office Word (Windows COM - %100 Aslına Uygunluk)
        if cls.is_ms_word_available():
            try:
                if progress_callback:
                    progress_callback(2, 3, "Microsoft Word motoru ile PDF oluşturuluyor...")
                res = cls._convert_via_ms_word(str(in_p), output_path)
                if os.path.exists(res):
                    # UYAP 10 MB Dosya Boyutu Koruması
                    try:
                        from core.redactor import PDFRedactor
                        PDFRedactor.optimize_for_uyap_if_needed(output_path)
                    except Exception:
                        pass

                    if progress_callback:
                        progress_callback(3, 3, "Dönüştürme başarıyla tamamlandı.")
                    return output_path
            except Exception:
                # MS Word başarısız olursa LibreOffice veya ReportLab fallback'e geç
                pass

        # 2. ÖNCELİK: LibreOffice Headless (Tam aslına uygun mizanpaj)
        soffice_bin = cls.find_libreoffice()
        if soffice_bin:
            try:
                temp_dir = tempfile.mkdtemp(prefix="paralegal_word_")
                if progress_callback:
                    progress_callback(2, 3, "LibreOffice motoru ile PDF çiziliyor...")

                cmd = [
                    soffice_bin,
                    "--headless",
                    "--convert-to",
                    "pdf",
                    str(in_p),
                    "--outdir",
                    temp_dir,
                ]

                creationflags = 0
                if sys.platform == "win32":
                    creationflags = getattr(subprocess, "CREATE_NO_WINDOW", 0x08000000)

                proc = subprocess.run(
                    cmd,
                    capture_output=True,
                    text=True,
                    timeout=90,
                    creationflags=creationflags,
                )

                gen_pdf = os.path.join(temp_dir, in_p.stem + ".pdf")
                if os.path.exists(gen_pdf):
                    shutil.move(gen_pdf, output_path)
                    try:
                        shutil.rmtree(temp_dir, ignore_errors=True)
                    except Exception:
                        pass

                    # UYAP 10 MB Dosya Boyutu Koruması
                    try:
                        from core.redactor import PDFRedactor
                        PDFRedactor.optimize_for_uyap_if_needed(output_path)
                    except Exception:
                        pass

                    if progress_callback:
                        progress_callback(3, 3, "Dönüştürme başarıyla tamamlandı.")
                    return output_path

            except Exception:
                # LibreOffice başarısız olursa Python fallback'e geç
                pass

        # 3. ÖNCELİK: Saf Python Fallback (python-docx + ReportLab) - Yalnızca .docx için
        if HAS_PYTHON_DOCX and HAS_REPORTLAB and in_p.suffix.lower() == ".docx":
            if progress_callback:
                progress_callback(2, 3, "Yerel Python kurtarma motoru ile PDF oluşturuluyor...")
            return cls._convert_via_reportlab(str(in_p), output_path)

        raise RuntimeError(
            "Word dönüştürme gerçekleştirilemedi. "
            "Lütfen sisteminizde Microsoft Word, LibreOffice veya gerekli Python paketlerinin (python-docx, reportlab) kurulu olduğunu doğrulayın."
        )

    @classmethod
    def _convert_via_reportlab(cls, input_path: str, output_path: str) -> str:
        """
        python-docx ve ReportLab kullanarak LibreOffice bağımlılığı olmadan Word'ü PDF'e çevirir.
        """
        doc = docx.Document(input_path)
        pdf = SimpleDocTemplate(
            output_path,
            pagesize=A4,
            leftMargin=54,  # ~1.9 cm standart adli kenar boşluğu
            rightMargin=54,
            topMargin=54,
            bottomMargin=54,
        )

        styles = getSampleStyleSheet()
        normal_style = styles["Normal"]
        normal_style.fontSize = 11
        normal_style.leading = 14
        normal_style.textColor = colors.HexColor("#121316")

        h1_style = ParagraphStyle(
            "Heading1_Custom",
            parent=styles["Heading1"],
            fontSize=14,
            leading=18,
            fontName="Helvetica-Bold",
            textColor=colors.HexColor("#121316"),
            spaceAfter=8,
        )

        elements = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                elements.append(Spacer(1, 8))
                continue

            # Başlık tespiti
            if p.style and "Heading" in p.style.name:
                elements.append(Paragraph(text, h1_style))
            else:
                elements.append(Paragraph(text, normal_style))
                elements.append(Spacer(1, 4))

        # Tabloları aktar
        for tbl in doc.tables:
            table_data = []
            for row in tbl.rows:
                row_data = [Paragraph(cell.text.strip(), normal_style) for cell in row.cells]
                table_data.append(row_data)

            if table_data:
                col_count = len(table_data[0])
                available_w = A4[0] - 108
                col_w = available_w / max(1, col_count)
                t = Table(table_data, colWidths=[col_w] * col_count)
                t.setStyle(
                    TableStyle([
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#121316")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("TOPPADDING", (0, 0), (-1, -1), 4),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                        ("LEFTPADDING", (0, 0), (-1, -1), 4),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                    ])
                )
                elements.append(t)
                elements.append(Spacer(1, 8))

        if not elements:
            elements.append(Paragraph("Boş Belge", normal_style))

        pdf.build(elements)

        # UYAP 10 MB Dosya Boyutu Koruması
        try:
            from core.redactor import PDFRedactor
            PDFRedactor.optimize_for_uyap_if_needed(output_path)
        except Exception:
            pass

        return output_path
